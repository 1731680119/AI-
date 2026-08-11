"""文件解析：图片 / PDF / Word / Excel / PowerPoint / 文本"""
import base64
import os
import uuid

import fitz  # PyMuPDF
import docx
import openpyxl
from pptx import Presentation
from PIL import Image, ImageOps
from io import BytesIO

from paths import UPLOAD_DIR as APP_UPLOAD_DIR


UPLOAD_DIR = str(APP_UPLOAD_DIR)

IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp"}
SHEET_EXTS = {".xlsx", ".xlsm"}
SLIDE_EXTS = {".pptx"}
# 纯文本按扩展名认，认不出来的二进制不硬读，避免给模型灌一堆乱码。
TEXT_EXTS = {".txt", ".md", ".markdown", ".rst", ".tex", ".csv", ".tsv", ".log",
             ".json", ".jsonl", ".xml", ".yaml", ".yml", ".toml", ".ini", ".cfg", ".properties",
             ".html", ".htm", ".css", ".scss", ".less", ".vue", ".svelte",
             ".py", ".js", ".mjs", ".cjs", ".ts", ".tsx", ".jsx",
             ".java", ".kt", ".scala", ".groovy", ".c", ".h", ".cpp", ".hpp", ".cc", ".cs",
             ".go", ".rs", ".rb", ".php", ".swift", ".m", ".mm", ".dart", ".lua", ".pl", ".r",
             ".sql", ".graphql", ".proto", ".sh", ".bash", ".zsh", ".bat", ".ps1", ".dockerfile"}
# 没有扩展名但按惯例是文本的文件名。
TEXT_NAMES = {"dockerfile", "makefile", "readme", "license", "changelog", ".gitignore",
              ".gitattributes", ".editorconfig"}

MAX_IMAGE_SIDE = 1568  # Claude 推荐最大边长
MAX_TEXT_CHARS = 200_000  # 单个文本附件的上限，防止一个大日志把上下文顶满


def save_upload(filename: str, data: bytes) -> dict:
    """保存上传文件，返回附件元信息"""
    os.makedirs(UPLOAD_DIR, exist_ok=True)
    ext = os.path.splitext(filename)[1].lower()
    fid = str(uuid.uuid4())
    stored_name = f"{fid}{ext}"
    path = os.path.join(UPLOAD_DIR, stored_name)
    with open(path, "wb") as f:
        f.write(data)

    kind = "file"
    preview = None
    text_content = None

    base = os.path.basename(filename).lower()
    if ext in IMAGE_EXTS:
        kind = "image"
        preview = _image_to_data_url(data, ext)
    elif ext == ".pdf":
        kind = "pdf"
        text_content = _parse_pdf(data)
    elif ext == ".docx":
        kind = "docx"
        text_content = _parse_docx(data)
    elif ext == ".doc":
        # 老版 .doc 是二进制格式，python-docx 读不了，直接说清楚而不是抛一段乱码。
        kind = "docx"
        text_content = "[旧版 .doc 格式无法解析，请另存为 .docx 后重新上传]"
    elif ext in SHEET_EXTS:
        kind = "sheet"
        text_content = _parse_xlsx(data)
    elif ext in SLIDE_EXTS:
        kind = "slide"
        text_content = _parse_pptx(data)
    elif ext == ".ppt":
        kind = "slide"
        text_content = "[旧版 .ppt 格式无法解析，请另存为 .pptx 后重新上传]"
    elif ext == ".xls":
        kind = "sheet"
        text_content = "[旧版 .xls 格式无法解析，请另存为 .xlsx 后重新上传]"
    elif ext in TEXT_EXTS or base in TEXT_NAMES or (not ext and base in TEXT_NAMES):
        kind = "text"
        text_content = _decode_text(data)
    else:
        # 认不出来的类型不猜，留个说明让模型知道自己拿不到内容。
        text_content = f"[不支持的文件类型 {ext or '(无扩展名)'}，未能提取文本]"

    if text_content and len(text_content) > MAX_TEXT_CHARS:
        text_content = text_content[:MAX_TEXT_CHARS] + f"\n[已截断，仅保留前 {MAX_TEXT_CHARS} 字]"

    return {
        "id": fid,
        "name": filename,
        "kind": kind,
        "stored_name": stored_name,
        "size": len(data),
        "preview": preview,          # 图片 data URL（缩略）
        "text_content": text_content # 文档解析出的文本
    }


def _image_to_data_url(data: bytes, ext: str) -> str:
    """压缩图片并转 data URL"""
    try:
        img = Image.open(BytesIO(data))
        img = ImageOps.exif_transpose(img)
        if max(img.size) > MAX_IMAGE_SIDE:
            img.thumbnail((MAX_IMAGE_SIDE, MAX_IMAGE_SIDE))
        buf = BytesIO()
        if img.mode in ("RGBA", "P"):
            img = img.convert("RGB")
        img.save(buf, format="JPEG", quality=85)
        b64 = base64.b64encode(buf.getvalue()).decode()
        return f"data:image/jpeg;base64,{b64}"
    except Exception:
        b64 = base64.b64encode(data).decode()
        mime = f"image/{ext.lstrip('.')}"
        return f"data:{mime};base64,{b64}"


def _parse_pdf(data: bytes, max_pages: int = 50) -> str:
    try:
        doc = fitz.open(stream=data, filetype="pdf")
        parts = []
        for i, page in enumerate(doc):
            if i >= max_pages:
                parts.append(f"\n[已截断，仅解析前 {max_pages} 页]")
                break
            parts.append(page.get_text())
        doc.close()
        return "\n".join(parts).strip()
    except Exception as e:
        return f"[PDF 解析失败: {e}]"


def _parse_docx(data: bytes) -> str:
    try:
        d = docx.Document(BytesIO(data))
        parts = [p.text for p in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                parts.append("\t".join(c.text for c in row.cells))
        return "\n".join(parts).strip()
    except Exception as e:
        return f"[Word 解析失败: {e}]"


def _decode_text(data: bytes) -> str:
    """按 UTF-8 / UTF-16 / GBK 依次尝试，都不成再带替换字符兜底。"""
    for encoding in ("utf-8-sig", "utf-16", "gbk"):
        try:
            return data.decode(encoding)
        except (UnicodeDecodeError, UnicodeError):
            continue
    return data.decode("utf-8", errors="replace")


def _parse_xlsx(data: bytes, max_rows: int = 500) -> str:
    """按工作表导出成 TSV。只取有值的区域，空表也标出来。"""
    try:
        # data_only 取公式算出的值；只读模式省内存。
        wb = openpyxl.load_workbook(BytesIO(data), data_only=True, read_only=True)
    except Exception as e:
        return f"[Excel 解析失败: {e}]"
    parts: list[str] = []
    try:
        for sheet in wb.worksheets:
            parts.append(f"## 工作表：{sheet.title}")
            count = 0
            for row in sheet.iter_rows(values_only=True):
                if count >= max_rows:
                    parts.append(f"[已截断，仅解析前 {max_rows} 行]")
                    break
                cells = ["" if v is None else str(v) for v in row]
                # 整行为空的行不占篇幅。
                if any(cell.strip() for cell in cells):
                    parts.append("\t".join(cells).rstrip())
                count += 1
            if count == 0:
                parts.append("(空工作表)")
            parts.append("")
    finally:
        wb.close()
    return "\n".join(parts).strip()


def _parse_pptx(data: bytes) -> str:
    """逐页导出文本框和表格内容，并保留备注。"""
    try:
        prs = Presentation(BytesIO(data))
    except Exception as e:
        return f"[PowerPoint 解析失败: {e}]"
    parts: list[str] = []
    for index, slide in enumerate(prs.slides, start=1):
        parts.append(f"## 第 {index} 页")
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = shape.text_frame.text.strip()
                if text:
                    parts.append(text)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    parts.append("\t".join(c.text for c in row.cells))
        if slide.has_notes_slide:
            notes = (slide.notes_slide.notes_text_frame.text or "").strip()
            if notes:
                parts.append(f"[备注] {notes}")
        parts.append("")
    return "\n".join(parts).strip()
